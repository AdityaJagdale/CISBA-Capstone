import os
import pdfplumber
from docx import Document as DocxDocument
from pptx import Presentation
import csv
from io import BytesIO, StringIO
import openpyxl
from langchain.text_splitter import CharacterTextSplitter
from langchain.docstore.document import Document

from helpers.genai import generate_answer


# Function to process files and generate answer
def process_files(uploaded_files, search_query, grounding_data):
    document_content = grounding_data

    try:
        for uploaded_file in uploaded_files:
            file_extension = os.path.splitext(uploaded_file.name)[1].lower()

            if file_extension == '.txt':
                document_content += uploaded_file.read().decode('utf-8') + "\n"

            elif file_extension == '.pdf':
                file_content = uploaded_file.read()
                with pdfplumber.open(BytesIO(file_content)) as pdf:
                    for page in pdf.pages:
                        text = page.extract_text()
                        if text:
                            document_content += text + "\n"

            elif file_extension in ['.pptx', '.ppt']:
                file_content = uploaded_file.read()
                prs = Presentation(BytesIO(file_content))
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            document_content += shape.text + "\n"

            elif file_extension in ['.docx', '.doc']:
                file_content = uploaded_file.read()
                doc = DocxDocument(BytesIO(file_content))
                for para in doc.paragraphs:
                    document_content += para.text + "\n"

            elif file_extension == '.csv':
                file_content = uploaded_file.read().decode('utf-8')
                reader = csv.reader(StringIO(file_content))
                for row in reader:
                    document_content += " ".join(row) + "\n"

            elif file_extension == '.xlsx':
                file_content = uploaded_file.read()
                workbook = openpyxl.load_workbook(BytesIO(file_content), data_only=True)
                for sheet in workbook.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        row_text = " ".join([str(cell) for cell in row if cell is not None])
                        document_content += row_text + "\n"

            else:
                return "Unsupported file type"

        if not document_content.strip():
            return "No text found in the document"

        # Split the document content into smaller chunks
        text_splitter = CharacterTextSplitter(chunk_size=800, chunk_overlap=100)
        texts = text_splitter.split_text(document_content.strip())

        # Create Document instances for each text chunk
        documents = [Document(page_content=text) for text in texts]

        context = document_content.strip()
        return generate_answer(context, search_query)

    except Exception as e:
        return f"Failed to load the document: {str(e)}"